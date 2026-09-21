import docx

doc = docx.Document('Реферат_Спорт_и_нормы.docx')

words = 0
for para in doc.paragraphs:
    words += len(para.text.split())

print(f"Total words: {words}")
print(f"Total paragraphs: {len(doc.paragraphs)}")
